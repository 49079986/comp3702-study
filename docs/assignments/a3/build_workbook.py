"""Build the study workbook from actual saved measurements. Not a submission script."""
import csv
import json
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "code"))
from solution import Solver


def load(group):
    return json.loads((ROOT / "results" / group / "summary.json").read_text(encoding="utf-8"))


def value(number, decimals=2):
    if number is None:
        return "未达到"
    return f"{number:,.{decimals}f}" if isinstance(number, float) else f"{number:,}"


def main():
    baseline = load("baseline")
    rates = load("learning_rates")
    comparison = load("comparison")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.72)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(11)
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.core_properties.title = "COMP3702 A3 学习工作稿"
    doc.core_properties.author = ""
    md = []

    def p(text, style=None):
        para = doc.add_paragraph(text, style)
        md.append(text + "\n")
        return para

    def heading(text, level=1):
        doc.add_heading(text, level)
        md.append("#" * (level + 1) + " " + text + "\n")

    def blank(lines=2):
        for _ in range(lines):
            para = doc.add_paragraph(" ")
            para.paragraph_format.space_after = Pt(3)
        md.append("[独立作答区域]\n")

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = str(text)
        for row in rows:
            for cell, text in zip(t.add_row().cells, row):
                cell.text = str(text)
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + side)
                    el.set(qn("w:w"), "65")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                props.append(margins)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "DFEBF3" if i == 0 else "FFFFFF")
                props.append(shading)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), "D9D9D9")
                    borders.append(border)
                props.append(borders)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.05
                    if j > 0 and len(str(cell.text)) < 24:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size = Pt(9)
                        run.font.bold = i == 0
            trpr = row._tr.get_or_add_trPr()
            cannot_split = OxmlElement("w:cantSplit")
            trpr.append(cannot_split)
        header_property = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(header_property)
        md.append("| " + " | ".join(headers) + " |")
        md.append("| " + " | ".join("---" for _ in headers) + " |")
        md.extend("| " + " | ".join(str(x) for x in row) + " |" for row in rows)
        md.append("")
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = Pt(3)
        spacer.add_run(" ").font.size = Pt(3)

    def image(relative, width=6.4):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(ROOT / relative), width=Inches(width))
        md.append(f"![学习曲线]({relative})\n")

    def next_page():
        doc.add_page_break()
        md.append("\n---\n")

    p("COMP3702 A3 学习工作稿", "Title")
    p("CrystalRover 强化学习  代码复核与实验记录", "Subtitle")
    p("这份工作稿供你运行参考实现、核对数据，再独立重写。允许 AI 的部分附代码对应说明和实测图表；禁止 AI 的小问仅保留题目和作答区域。")
    p("姓名：________________  学号：________________  邮箱：________________")
    p("作业标注截止时间：2026 年 10 月 30 日 16:00。代码与报告各占 50 分。")
    heading("Q1 Q learning 与 Value Iteration")
    p("本题共 5 分。Q1a 与 Q1b 均禁止 AI。")
    heading("Q1a 相似点 2点5分", 2)
    p("描述标准表格型 Q-learning 与标准表格型 Value Iteration 的两个关键相似点，答案不超过五行。")
    blank(5)
    heading("Q1b 不同点 2点5分", 2)
    p("描述两个关键不同点，答案不超过五行。不要换一种表述重复同一个核心点。")
    blank(5)
    p("Q1 的答案由你独立填写。")

    next_page()
    heading("Q2 实现与学习率实验")
    heading("Q2a 实现说明 3分 允许AI但需明确验证", 2)
    p("参考 QL 使用 epsilon-greedy。每个 episode 开始时根据已用训练步数计算 epsilon，训练中的贪心并列动作随机打破并列；评估方法使用确定性的贪心选择。")
    table(["参数", "参考实现设置"], [
        ["QL 学习率", "固定 alpha = 0.2"],
        ["epsilon", "从 0.25 线性降到 0.02"],
        ["衰减区间", "训练步数预算的前 80%"],
        ["初始化", "Q 表按需创建，初始 Q 值为 0"],
        ["实验候选学习率", "0.1  0.2  0.4，均为固定值"],
    ], [2.0, 4.6])
    p("这些值用于演示与复现，不宣称已经最优。最终采用值、替代方案及选择依据，请根据你的实际实验与最终实现补充：________________。")
    p("自动检查已核对单步终止更新、合法动作、训练状态链、有限特征与评估不改权重。你的验证记录：检查命令、手算样本、实际观察和发现的问题：________________。")
    heading("Q2b 三个固定学习率 4分 允许AI", 2)
    p("L1，环境与探索种子 1914。三次独立训练只改变 alpha，均使用官方步数及累计奖励预算；最终策略评估 1000 次。图中为训练回报的 50-episode 移动平均。")
    image("results/learning_rates/learning_curves.png", 5.85)
    table(["alpha", "训练episode", "评估平均回报", "成功次数"], [
        [f"{r['alpha']:g}", value(r["training_episodes"]),
         value(r["evaluation_mean_reward"]), f"{r['solved_trials']}/1000"] for r in rates
    ], [0.7, 1.5, 2.2, 2.2])

    next_page()
    heading("Q2c 独立作答")
    p("本小问 6 分，禁止 AI。结合自己的图表，讨论学习率对学习速度、稳定性和最终观测表现的影响，解释意外结果及实验限制。")
    blank(2)
    heading("Q3a 特征设计 4分 允许AI但需明确验证")
    p("FQ 使用 21 维固定长度向量，Q 值为学习权重与特征的点积。权重从全 0 开始，以固定 alpha=0.03 做样本半梯度更新。以下为归一化前的特征定义。")
    table(["编号", "特征", "计算与作用"], [
        ["0至7", "bias 类型 方向", "常数、walk/boost/jump 指示、左/右/上/下指示，表示动作类别与朝向。"],
        ["8至9", "目标距离", "当前格、前方一格的地图 BFS 距离，除以行数加列数。"],
        ["10至12", "朝目标的进展", "当前距离减前方一格距离；boost 另看两格与四格，除以同一尺度。"],
        ["13至14", "岩浆位置", "正前方是否岩浆，以及两个侧邻格的岩浆指示平均值。"],
        ["15至16", "岩石与陨石坑", "正前方格的地形指示，描述方向上的障碍和坑。"],
        ["17至18", "收集进度", "尚缺样本数除以最低要求，以及是否已满足最低数量。"],
        ["19至20", "目标在前方", "前方一格是否有未收集水晶；数量足够后是否是发射点。"],
    ], [0.65, 1.7, 4.25])
    p("目标在样本不足时是最近未收集水晶，足够后是最近发射点。BFS 仅查看公开地图，将岩石和岩浆视为不可通行，未使用隐藏概率或成本。直线方向查询遇岩石、岩浆或坑时停止，不代表真实随机转移概率。")
    p("原始向量统一除以 max(1, L2范数)，保持数值尺度有限。因此归一化后的 bias 不一定是 1。局限是最近目标不代表全程最优目标组合，距离忽略风暴、随机移动和真实动作成本，局部地形也不能描述所有远处风险。")
    p("验证请实际提取若干状态与动作的向量，手算对应距离与地形，并用官方特征检查核对长度。你的检查记录与修订：________________。")

    next_page()
    heading("Q3b 四个关卡的学习权重 2分 允许AI")
    p("使用 baseline 中各关卡官方 FQ 预算训练结束后的原始权重。列名与代码 FEATURE_NAMES 一一对应；所有特征都经过前述向量归一化。")
    weight_sets = {}
    for level in ("L1", "L2", "L3", "L4"):
        with (ROOT / "results" / "baseline" / f"{level}_fq_weights.csv").open(encoding="utf-8-sig") as stream:
            weight_sets[level] = {r["feature"]: float(r["weight"]) for r in csv.DictReader(stream)}
    weight_rows = [[name] + [f"{weight_sets[level][name]:.4f}" for level in ("L1", "L2", "L3", "L4")]
                   for name in Solver.FEATURE_NAMES]
    table(["Feature", "L1", "L2", "L3", "L4"], weight_rows, [2.65, 0.98, 0.98, 0.98, 0.98])
    heading("Q3c 权重解读 6分 禁止AI", 2)
    p("结合特征定义讨论权重的符号和大小、是否符合你对良好策略的预期，以及关卡之间的差异。")
    blank(2)

    next_page()
    heading("Q4a QL 与 FQ 实验比较 6分 允许AI")
    p("选择 L1 和 L3：L1 为无岩浆走廊地图；L3 包含岩浆、陨石坑且需三个样本，以覆盖不同地形与收集要求。")
    p("两种算法每关共用 30000 步上限和相同累计奖励下限，使用相同评估种子、1000 次最终评估与同一个阈值。回报区间取两种官方区间的包络，阈值为下限加区间宽度的 75%。学习率分别为 QL 0.2、FQ 0.03；epsilon 设置相同。")
    table(["关卡", "累计奖励下限", "学习阈值", "评估截断"], [
        [level, value(next(r for r in comparison if r["level"] == level)["reward_budget"]),
         value(next(r for r in comparison if r["level"] == level)["learning_threshold"], 3),
         value(next(r for r in comparison if r["level"] == level)["evaluation_reward_floor"])]
        for level in ("L1", "L3")
    ], [0.7, 2.0, 2.0, 1.9])
    image("results/comparison/learning_curves.png", 5.8)
    table(["运行", "达标episode", "达标步数", "平均回报", "成功次数"], [
        [r["run"], value(r["learned_at_episodes"]), value(r["learned_at_steps"]),
         value(r["evaluation_mean_reward"]), f"{r['solved_trials']}/1000"] for r in comparison
    ], [0.9, 1.3, 1.3, 1.5, 1.6])
    p("阈值来自每 1/25 训练步数预算的一次 50-episode 贪心评估。未达到保留为未达到；训练回报曲线与贪心评估不是同一统计量。")

    next_page()
    heading("Q4b 概念预测 6分 禁止AI")
    p("预测哪种算法学习更快、哪种最终表现更好，并说明理由。讨论表格型和特征型表示的优点、局限、泛化和近似误差。")
    blank(5)
    heading("Q4c 结果解读与结论 8分 禁止AI")
    p("结合自己的图表比较学习轨迹、速度、稳定性与最终质量；核对概念预测，讨论两个关卡的差异，并以实验支持解释。")
    blank(5)
    heading("附录 AI声明与参考来源")
    p("AI 使用记录草稿：OpenAI Codex 于 2026 年 10 月 26 日协助生成学习参考代码、实验脚本、实测图表及本工作稿允许 AI 的部分。模型版本：按你所用界面实际信息填写。")
    p("请按实际情况补充最终保留的内容、你做过的验证与修订。禁止 AI 的答案需由你独立完成，不要声明尚未做过的验证。")
    p("参考来源：COMP3702 Assignment 3 CrystalRover Reinforcement Learning，2026，课程提供的作业说明。")
    p("官方支持代码与接口：https://github.com/comp3702-2026/2026-Assignment-3-Support-Code")
    p("支持代码版本：0e498b0adbe94b0e6b084945e575de3df4988d0e。原始实验数据保存在 results 的三个分组中。")

    for border in list(doc.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.save(ROOT / "A3_报告学习工作稿.docx")
    (ROOT / "报告学习工作稿.md").write_text("\n".join(md), encoding="utf-8")
    # A separate factual record of checks and the official-budget baseline.
    validation = ["# 验证记录", "", "以下为本机实测记录，不是 Gradescope 得分。", "",
                  "## 协议检查", "", "```text",
                  (ROOT / "tmp" / "protocol_checks.txt").read_text(encoding="utf-8-sig"), "```", "",
                  "## 官方预算下的最终评估", "",
                  "| 运行 | 训练步数 | episode数 | 平均回报 | 成功次数 | 特征数 | 泛化平均回报 |",
                  "| --- | --- | --- | --- | --- | --- | --- |"]
    for r in baseline:
        generalisation = "不适用" if r["generalisation_mean_reward"] is None else value(r["generalisation_mean_reward"])
        validation.append(f"| {r['run']} | {r['training_steps']} | {r['training_episodes']} | "
                          f"{r['evaluation_mean_reward']:.2f} | {r['solved_trials']}/1000 | "
                          f"{r['feature_count'] or '不适用'} | {generalisation} |")
    validation.extend(["", "每个算法与关卡仅展示一个官方种子；不是多次独立重复的统计结论。计时保存在 JSON 中，受本机负载影响。",
                       "", "所有 summary 记录对应的 solution.py SHA256：`" + baseline[0]["solver_sha256"] + "`。"])
    (ROOT / "验证记录.md").write_text("\n".join(validation), encoding="utf-8")
    print("Created workbook, Markdown practice draft, and verification record.")


if __name__ == "__main__":
    main()
